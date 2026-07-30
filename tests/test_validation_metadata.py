import datetime
import sys
import tempfile
import unittest
from enum import Enum
from pathlib import Path
from typing import Annotated, Any, Dict, List, Optional
from typing import get_args, get_origin

import openpyxl
import pandas as pd
from pydantic import BaseModel, StringConstraints, ValidationError

repo_root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(repo_root / "src"))

from pegasus.main import cross_validate_list_matrix
from pegasus.validation.metadata_validation import (
    PegMetadataValidation,
    _format_excel_datetime,
)
from pegasus.schema.peg_metadata_schema.metadata_basic_schema import (
    DatasetDescription,
    GenomicIdentifier,
)
from pegasus.schema.peg_metadata_schema.metadata_evidence_schema import Evidence
from pegasus.schema.peg_metadata_schema.metadata_integration_schme import Integration
from pegasus.schema.peg_metadata_schema.metadata_method_schema import Method
from pegasus.schema.peg_metadata_schema.metadata_source_schema import Source


def _fallback_value(annotation: Any) -> Any:
    ann = annotation
    origin = get_origin(ann)
    args = get_args(ann)
    if origin is not None and str(origin).endswith("Annotated"):
        ann = args[0]
        origin = get_origin(ann)
        args = get_args(ann)
    if origin is not None and str(origin).endswith("Union"):
        non_none = [a for a in args if a is not type(None)]
        if len(non_none) == 1:
            ann = non_none[0]
            origin = get_origin(ann)
            args = get_args(ann)
    if origin is not None and str(origin).endswith("Literal"):
        return args[0]
    try:
        if isinstance(ann, type) and issubclass(ann, Enum):
            return list(ann)[0].value
    except Exception:
        pass
    if ann is bool:
        return True
    if ann is int:
        return 1
    if ann is float:
        return 0.1
    if ann is str:
        return "value"
    name = getattr(ann, "__name__", "")
    if name == "HttpUrl":
        return "https://example.com"
    return "value"


def _example_row(model: type) -> Dict[str, Any]:
    row: Dict[str, Any] = {}
    for field_name, field in model.model_fields.items():
        extra = field.json_schema_extra or {}
        header = extra.get("header", field_name)
        if "example" in extra:
            value = extra["example"]
        else:
            value = _fallback_value(field.annotation)
        row[str(header)] = value
    if model is Evidence:
        row["evidence_category"] = "Molecular QTL"
        row["evidence_category_abbreviation"] = "Molecular QTL"
        row["variant_or_gene_centric"] = "variant-centric"
    if model is Integration:
        row["author_conclusion"] = False
    return row


def _build_sheet_df(model: type, overrides: List[Dict[str, Any]]) -> pd.DataFrame:
    base = _example_row(model)
    rows = []
    for override in overrides:
        row = dict(base)
        row.update(override)
        rows.append(row)
    return pd.DataFrame(rows)

def _empty_sheet_df(model: type) -> pd.DataFrame:
    base = _example_row(model)
    return pd.DataFrame(columns=base.keys())


def _write_metadata_excel(
    path: Path,
    evidence_rows: List[Dict[str, Any]],
    integration_rows: List[Dict[str, Any]],
    source_rows: Optional[List[Dict[str, Any]]] = None,
    method_rows: Optional[List[Dict[str, Any]]] = None,
    dataset_rows: Optional[List[Dict[str, Any]]] = None,
    genomic_rows: Optional[List[Dict[str, Any]]] = None,
    omit_sheets: Optional[List[str]] = None,
    drop_columns: Optional[Dict[str, List[str]]] = None,
) -> None:
    normalized_omit = {name.lower().strip() for name in (omit_sheets or [])}
    normalized_drop = {
        name.lower().strip(): cols for name, cols in (drop_columns or {}).items()
    }

    dataset_df = _build_sheet_df(DatasetDescription, dataset_rows or [{}])
    genomic_df = _build_sheet_df(GenomicIdentifier, genomic_rows or [{}])
    evidence_df = _build_sheet_df(Evidence, evidence_rows)
    integration_df = (
        _empty_sheet_df(Integration)
        if integration_rows == []
        else _build_sheet_df(Integration, integration_rows)
    )
    source_df = _build_sheet_df(Source, source_rows or [{}])
    method_df = _build_sheet_df(Method, method_rows or [{}])

    if normalized_drop.get("datasetdescription"):
        dataset_df = dataset_df.drop(columns=normalized_drop["datasetdescription"], errors="ignore")
    if normalized_drop.get("genomicidentifier"):
        genomic_df = genomic_df.drop(columns=normalized_drop["genomicidentifier"], errors="ignore")
    if normalized_drop.get("evidence"):
        evidence_df = evidence_df.drop(columns=normalized_drop["evidence"], errors="ignore")
    if normalized_drop.get("integration"):
        integration_df = integration_df.drop(columns=normalized_drop["integration"], errors="ignore")
    if normalized_drop.get("source"):
        source_df = source_df.drop(columns=normalized_drop["source"], errors="ignore")
    if normalized_drop.get("method"):
        method_df = method_df.drop(columns=normalized_drop["method"], errors="ignore")

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        if "datasetdescription" not in normalized_omit:
            dataset_df.to_excel(writer, sheet_name="DatasetDescription", index=False, startrow=1)
        if "genomicidentifier" not in normalized_omit:
            genomic_df.to_excel(writer, sheet_name="GenomicIdentifier", index=False, startrow=1)
        if "evidence" not in normalized_omit:
            evidence_df.to_excel(writer, sheet_name="Evidence", index=False, startrow=1)
        if "integration" not in normalized_omit:
            integration_df.to_excel(writer, sheet_name="Integration", index=False, startrow=1)
        if "source" not in normalized_omit:
            source_df.to_excel(writer, sheet_name="Source", index=False, startrow=1)
        if "method" not in normalized_omit:
            method_df.to_excel(writer, sheet_name="Method", index=False, startrow=1)


def _has_step(results: List[Dict[str, Any]], step: str) -> bool:
    return any(result.get("step") == step for result in results)

def _has_type(results: List[Dict[str, Any]], kind: str) -> bool:
    return any(result.get("type") == kind for result in results)

def _valid_evidence_rows(
    count: int,
    source_tag: str = "source_ok",
    method_tag: str = "method_ok",
) -> List[Dict[str, Any]]:
    return [
        {
            "column_header": f"EV{i}",
            "column_description": f"Evidence {i}",
            "evidence_category": "Molecular QTL",
            "evidence_category_abbreviation": "QTL",
            "variant_or_gene_centric": "variant-centric",
            "source_tag": source_tag,
            "method_tag": method_tag,
        }
        for i in range(1, count + 1)
    ]

def _valid_integration_rows(
    count: int,
    author_conclusion_index: int = 0,
    method_tag: str = "method_ok",
) -> List[Dict[str, Any]]:
    return [
        {
            "integration_tag": f"int{i}",
            "column_header": f"INT{i}",
            "column_description": f"Integration {i}",
            "author_conclusion": i - 1 == author_conclusion_index,
            "method_tag": method_tag,
        }
        for i in range(1, count + 1)
    ]


class TestMetadataValidation(unittest.TestCase):
    def test_evidence_requires_more_than_two_valid_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(2),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Evidence - Row Count"))

    def test_integration_requires_more_than_one_valid_row(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(2, author_conclusion_index=1),
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Integration - Row Count"))

    def test_author_conclusion_requires_exactly_one_true(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            rows = _valid_integration_rows(3, author_conclusion_index=1)
            rows[2]["author_conclusion"] = True
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=rows,
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Integration - Author Conclusion"))

    def test_tag_cross_validation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4, source_tag="source_missing", method_tag="method_missing"),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1, method_tag="method_missing2"),
                source_rows=[{}, {"source_tag": "source_ok"}],
                method_rows=[{}, {"method_tag": "method_ok"}],
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Evidence - Tag Reference"))
            self.assertTrue(_has_step(results, "Integration - Tag Reference"))

    def test_catalog_success(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                source_rows=[{}, {"source_tag": "source_ok"}],
                method_rows=[{}, {"method_tag": "method_ok"}],
                dataset_rows=[{}, {"trait_description": "Trait A"}],
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertFalse(_has_type(results, "error"))

    def test_catalog_missing_essential_tab(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                omit_sheets=["DatasetDescription"],
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Sheet Validation"))

    def test_catalog_missing_mandatory_columns_dataset(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                drop_columns={"DatasetDescription": ["trait_description"]},
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "DatasetDescription - Header Validation"))

    def test_catalog_missing_mandatory_columns_identifier(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                drop_columns={"GenomicIdentifier": ["genome_build"]},
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "GenomicIdentifier - Header Validation"))

    def test_catalog_missing_mandatory_columns_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                drop_columns={"Evidence": ["column_header"]},
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Evidence - Header Validation"))

    def test_catalog_missing_mandatory_columns_integration(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                drop_columns={"Integration": ["column_description"]},
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Integration - Header Validation"))

    def test_catalog_gwas_source_is_gwas_catalog(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            gwas_overrides = {
                "gwas_samples_description": None,
                "gwas_sample_size": None,
                "gwas_case_control_study": None,
                "gwas_sample_ancestry": None,
                "gwas_sample_ancestry_label": None,
            }
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                source_rows=[{}, {"source_tag": "source_ok"}],
                method_rows=[{}, {"method_tag": "method_ok"}],
                dataset_rows=[{}, {"gwas_source": "GCST123456", **gwas_overrides}],
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertFalse(_has_type(results, "error"))

    def test_catalog_gwas_source_is_not_gwas_catalog(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            gwas_overrides = {
                "gwas_samples_description": None,
                "gwas_sample_size": None,
                "gwas_case_control_study": None,
                "gwas_sample_ancestry": None,
                "gwas_sample_ancestry_label": None,
            }
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                dataset_rows=[{}, {"gwas_source": "PMID:1234", **gwas_overrides}],
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "DatasetDescription - Row Validation"))

    def test_catalog_evidence_category_not_in_list(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            rows = _valid_evidence_rows(4)
            rows[1]["evidence_category"] = "UNKNOWN"
            _write_metadata_excel(
                path,
                evidence_rows=rows,
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Evidence - Row Validation"))

    def test_catalog_evidence_category_only_one(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(2),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Evidence - Row Count"))

    def test_catalog_no_integration_row(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=[],
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Integration - Row Count"))

    def test_catalog_more_than_one_author_conclusion(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            rows = _valid_integration_rows(3, author_conclusion_index=1)
            rows[2]["author_conclusion"] = True
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=rows,
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Integration - Author Conclusion"))

    def test_catalog_no_author_conclusion(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            rows = _valid_integration_rows(3, author_conclusion_index=2)
            rows[1]["author_conclusion"] = False
            rows[2]["author_conclusion"] = False
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=rows,
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Integration - Author Conclusion"))

    def test_catalog_evidence_category_miss_source_tag(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4, source_tag="source_missing"),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                source_rows=[{}, {"source_tag": "source_ok"}],
                method_rows=[{}, {"method_tag": "method_ok"}],
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Evidence - Tag Reference"))

    def test_catalog_evidence_category_miss_method_tag(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(4),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1, method_tag="method_missing"),
                source_rows=[{}, {"source_tag": "source_ok"}],
                method_rows=[{}, {"method_tag": "method_ok"}],
            )
            results = PegMetadataValidation(path).validate_metadata()
            self.assertTrue(_has_step(results, "Integration - Tag Reference"))


class TestBlankColumnHeaderResidue(unittest.TestCase):
    """Residue rows that declare no column name must not reach cross-validation.

    A row with a blank column_header but other cells still populated survives the
    empty-row filter, so it used to contribute a None column name and crash
    cross-validation on sorted().

    In the workbook that exposed this, the populated cells held "#REF!" — broken
    Excel references. The fixture leaves them at their example values instead,
    because openpyxl writes a literal "#REF!" as an error-typed cell that pandas
    reads back as NaN, which would make the whole row empty and filtered. What
    matters is the shape: no column_header, everything else present.
    """

    _RESIDUE_ROW = {"column_header": None}

    def _write_workbook(self, path: Path) -> None:
        _write_metadata_excel(
            path,
            evidence_rows=_valid_evidence_rows(3) + [dict(self._RESIDUE_ROW)],
            integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
            source_rows=[{}, {"source_tag": "source_ok"}],
            method_rows=[{}, {"method_tag": "method_ok"}],
        )

    def test_blank_column_headers_are_excluded(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            self._write_workbook(path)

            validator = PegMetadataValidation(path)
            validator.validate_metadata()
            evidence_headers = validator.cross_check_column_names()["Evidence"]

            self.assertTrue(evidence_headers, "real evidence columns should survive")
            for name in evidence_headers:
                self.assertIsInstance(name, str)
                self.assertTrue(name.strip(), f"blank column header leaked through: {name!r}")

    def test_cross_validation_reports_instead_of_crashing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            metadata_file = tmp_path / "metadata_residue.xlsx"
            self._write_workbook(metadata_file)

            list_file = tmp_path / "list_residue.tsv"
            list_file.write_text(
                "PrimaryVariantID\tGeneSymbol\tGWAS\tINT_Combined_score\n"
                "chr1:100000:A:G\tVTI1A\tTRUE\tSTRONG\n",
                encoding="utf-8",
            )
            matrix_file = tmp_path / "matrix_residue.tsv"
            matrix_file.write_text(
                "PrimaryVariantID\tGeneSymbol\tGWAS_pvalue\tINT_score\n"
                "chr1:100000:A:G\tVTI1A\t4e-8\tSTRONG\n",
                encoding="utf-8",
            )

            results = cross_validate_list_matrix(list_file, matrix_file, metadata_file)

            # The metadata declares columns the matrix lacks, so this run does
            # report a mismatch. The point is that it reports rather than raising.
            self.assertIsInstance(results, list)
            self.assertTrue(_has_type(results, "error"))


def _write_companion_tsvs(tmp_path: Path, stem: str) -> tuple[Path, Path]:
    """Write a minimal list/matrix pair. Cross-validation only reads their headers."""
    list_file = tmp_path / f"list_{stem}.tsv"
    list_file.write_text(
        "PrimaryVariantID\tGeneSymbol\tGWAS\tINT_Combined_score\n"
        "chr1:100000:A:G\tVTI1A\tTRUE\tSTRONG\n",
        encoding="utf-8",
    )
    matrix_file = tmp_path / f"matrix_{stem}.tsv"
    matrix_file.write_text(
        "PrimaryVariantID\tGeneSymbol\tGWAS_pvalue\tINT_score\n"
        "chr1:100000:A:G\tVTI1A\t4e-8\tSTRONG\n",
        encoding="utf-8",
    )
    return list_file, matrix_file


class TestAuthorConclusionCountReporting(unittest.TestCase):
    """The metadata must declare exactly one author-conclusion row.

    return_author_conclusion_rows raises when it does not, so cross-validation
    has to catch that and report it rather than let it kill the CLI.
    """

    def _assert_reported(self, integration_rows: List[Dict[str, Any]]) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            metadata_file = tmp_path / "metadata_ac.xlsx"
            _write_metadata_excel(
                metadata_file,
                evidence_rows=_valid_evidence_rows(3),
                integration_rows=integration_rows,
                source_rows=[{}, {"source_tag": "source_ok"}],
                method_rows=[{}, {"method_tag": "method_ok"}],
            )
            list_file, matrix_file = _write_companion_tsvs(tmp_path, "ac")

            results = cross_validate_list_matrix(list_file, matrix_file, metadata_file)

            self.assertIsInstance(results, list)
            error_steps = [r.get("step", "") for r in results if r.get("type") == "error"]
            self.assertTrue(
                any("Author Conclusion Records" in step for step in error_steps),
                f"expected an author-conclusion error, got: {[r.get('step') for r in results]}",
            )

    def test_zero_author_conclusion_rows_is_reported(self) -> None:
        # overrides[0] is consumed as the example row, so flagging it leaves none.
        self._assert_reported(_valid_integration_rows(3, author_conclusion_index=0))

    def test_two_author_conclusion_rows_is_reported(self) -> None:
        rows = _valid_integration_rows(3, author_conclusion_index=1)
        rows[2]["author_conclusion"] = True
        self._assert_reported(rows)


class TestReportedRowNumbers(unittest.TestCase):
    """A reported row number must point at the row the curator sees in Excel."""

    def test_reported_row_matches_spreadsheet_row(self) -> None:
        bad_url = "definitely not a url"
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(3),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                source_rows=[{}, {"source_tag": "source_ok", "url": bad_url}, {}],
                method_rows=[{}, {"method_tag": "method_ok"}],
            )

            # Find where the bad value actually sits, rather than hard-coding it.
            sheet = openpyxl.load_workbook(path)["Source"]
            headers = [cell.value for cell in sheet[2]]
            url_col = headers.index("url") + 1
            expected_row = next(
                r for r in range(3, sheet.max_row + 1)
                if sheet.cell(r, url_col).value == bad_url
            )

            results = PegMetadataValidation(path).validate_metadata()
            source_errors = [
                r for r in results
                if r.get("step") == "Source - Row Validation" and r.get("type") == "error"
            ]
            self.assertEqual(len(source_errors), 1, f"expected one Source row error, got {results}")

            reported_rows = [detail["row"] for detail in source_errors[0]["details"]]
            self.assertEqual(reported_rows, [expected_row])


class _PatternModel(BaseModel):
    accession: Annotated[str, StringConstraints(pattern=r"^GCST\d+$")]


class TestErrorEnrichment(unittest.TestCase):
    """Enriched errors must not mislead about whitespace or dump whole rows."""

    @staticmethod
    def _errors_for(model: type, record: Dict[str, Any]) -> List[Dict[str, Any]]:
        try:
            model.model_validate(record)
        except ValidationError as exc:
            return PegMetadataValidation._deduplicate_errors(exc.errors())
        raise AssertionError("expected validation to fail")

    def _enrich(self, errors: List[Dict[str, Any]], field: str, examples: Dict[str, Any]) -> Dict[str, Any]:
        matching = [e for e in errors if list(e.get("loc") or [None])[0] == field]
        self.assertTrue(matching, f"no error reported for {field}: {errors}")
        return PegMetadataValidation._enrich_error(matching[0], examples)

    def test_no_whitespace_hint_for_enum_field(self) -> None:
        """"Molecular QTL" contains a space, so the hint would contradict itself."""
        record = {
            "column_header": "EV1",
            "column_description": "desc",
            "evidence_category": "Totally Wrong Category",
            "evidence_category_abbreviation": "QTL",
            "variant_or_gene_centric": "variant-centric",
        }
        errors = self._errors_for(Evidence, record)
        entry = self._enrich(errors, "evidence_category", {"evidence_category": "Molecular QTL"})

        self.assertNotIn("hint", entry)
        self.assertEqual(entry.get("expected_example"), "Molecular QTL")

    def test_whitespace_hint_kept_for_pattern_field(self) -> None:
        errors = self._errors_for(_PatternModel, {"accession": "GCST 000001"})
        entry = self._enrich(errors, "accession", {"accession": "GCST000001"})

        self.assertIn("hint", entry)
        self.assertIn("whitespace", entry["hint"])

    def test_missing_field_error_omits_the_row_dump(self) -> None:
        errors = self._errors_for(_PatternModel, {})
        entry = self._enrich(errors, "accession", {})

        self.assertNotIn("value", entry)


class TestExcelDateCoercion(unittest.TestCase):
    """Excel date cells arrive as datetimes and must not fail string fields."""

    def test_formats_midnight_as_plain_date(self) -> None:
        self.assertEqual(_format_excel_datetime(datetime.datetime(2020, 6, 6)), "2020-06-06")

    def test_formats_date_as_plain_date(self) -> None:
        self.assertEqual(_format_excel_datetime(datetime.date(2020, 6, 6)), "2020-06-06")

    def test_keeps_time_component_when_present(self) -> None:
        self.assertEqual(
            _format_excel_datetime(datetime.datetime(2020, 6, 6, 13, 45)),
            "2020-06-06T13:45:00",
        )

    def test_date_cell_validates_as_a_string_field(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "meta.xlsx"
            _write_metadata_excel(
                path,
                evidence_rows=_valid_evidence_rows(3),
                integration_rows=_valid_integration_rows(3, author_conclusion_index=1),
                source_rows=[{}, {"source_tag": "source_ok", "version": datetime.datetime(2020, 6, 6)}],
                method_rows=[{}, {"method_tag": "method_ok"}],
            )

            validator = PegMetadataValidation(path)
            results = validator.validate_metadata()

            versions = [r.get("version") for r in validator.sheet_data["Source"]["records"]]
            self.assertIn("2020-06-06", versions)

            source_errors = [
                r for r in results
                if r.get("step") == "Source - Row Validation" and r.get("type") == "error"
            ]
            self.assertEqual(source_errors, [], f"date cell should validate cleanly: {source_errors}")
