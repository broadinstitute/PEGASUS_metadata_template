import argparse
import sys
import tempfile
import unittest
from pathlib import Path
from typing import Optional

import openpyxl

repo_root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(repo_root / "src"))

from pegasus.main import handle_convert

EXISTING_CONTENT = b"IMPORTANT USER DATA\n"


def _args(
    output_path: Optional[Path],
    force: bool = False,
    input_path: Optional[Path] = None,
) -> argparse.Namespace:
    return argparse.Namespace(
        command="convert",
        conversion_type="schema-to-xlsx",
        input_path=input_path,
        output_path=output_path,
        force=force,
    )


class TestSchemaToXlsxOverwriteGuard(unittest.TestCase):
    """schema-to-xlsx has no input, so a lone positional is the destination.

    That makes it easy to name a file you meant to keep, so an existing file must
    survive unless --force is passed.
    """

    def test_refuses_to_overwrite_existing_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            target = Path(tmp_dir) / "my_metadata.xlsx"
            target.write_bytes(EXISTING_CONTENT)

            exit_code = handle_convert(_args(target))

            self.assertEqual(exit_code, 1)
            self.assertEqual(target.read_bytes(), EXISTING_CONTENT)

    def test_refuses_to_overwrite_when_named_as_input_path(self) -> None:
        """The real footgun: the destination arrives as input_path, not output_path."""
        with tempfile.TemporaryDirectory() as tmp_dir:
            target = Path(tmp_dir) / "my_metadata.xlsx"
            target.write_bytes(EXISTING_CONTENT)

            exit_code = handle_convert(_args(None, input_path=target))

            self.assertEqual(exit_code, 1)
            self.assertEqual(target.read_bytes(), EXISTING_CONTENT)

    def test_force_allows_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            target = Path(tmp_dir) / "my_metadata.xlsx"
            target.write_bytes(EXISTING_CONTENT)

            exit_code = handle_convert(_args(target, force=True))

            self.assertEqual(exit_code, 0)
            self.assertIn("DatasetDescription", openpyxl.load_workbook(target).sheetnames)

    def test_creates_a_new_file_without_force(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            target = Path(tmp_dir) / "fresh_template.xlsx"

            exit_code = handle_convert(_args(target))

            self.assertEqual(exit_code, 0)
            self.assertIn("DatasetDescription", openpyxl.load_workbook(target).sheetnames)

    def test_missing_destination_is_an_error(self) -> None:
        self.assertEqual(handle_convert(_args(None)), 1)


if __name__ == "__main__":
    unittest.main()
